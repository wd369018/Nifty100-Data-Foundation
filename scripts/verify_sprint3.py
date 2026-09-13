"""Sprint 3 verification harness.

Re-runs every Sprint 3 (Days 15-21) exit criterion against the live
database and shared outputs, mirroring `scripts/verify_sprint2.py`.

Run:  python scripts/verify_sprint3.py
"""

import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from src.screener.engine import (
    compute_composite_score,
    load_config,
    load_feature_frame,
    run_preset,
)

ROOT = Path(__file__).resolve().parents[1]

PASS = "PASS"
FAIL = "FAIL"


def check(name, passed, detail=""):
    status = PASS if passed else FAIL
    print(f"  [{status}] {name}" + (f" — {detail}" if detail else ""))
    return passed


def main():
    failures = []

    print("=" * 70)
    print("SPRINT 3 VERIFICATION (Days 15-21 — Screener + Peer Engine)")
    print("=" * 70)

    # ------------------------------------------------------------
    # 1. Screener: 6 presets each return 5-50 companies
    # ------------------------------------------------------------
    print("\n[1] Screener presets (5-50 companies each)")
    config = load_config()
    frame = compute_composite_score(load_feature_frame())
    counts = {}
    for key in config["presets"]:
        result = run_preset(frame, config, key)
        n = len(result["frame"])
        counts[key] = n
        name = result["name"]
        failures.append(
            not check(
                f"{name:<20s}",
                5 <= n <= 50,
                f"{n} companies",
            )
        )

    # ------------------------------------------------------------
    # 2. Quality Compounder top 5 are profitable compounders
    # ------------------------------------------------------------
    print("\n[2] Quality Compounder top 5 spot check")
    qc = run_preset(frame, config, "quality_compounder")["frame"].head(5)
    ok = True
    for _, row in qc.iterrows():
        valid = (
            row["return_on_equity_pct"] > 15
            and row["free_cash_flow_cr"] > 0
            and row["revenue_cagr_5yr"] > 10
            and (row["broad_sector"] == "Financials" or row["debt_to_equity"] < 1)
        )
        ok = ok and valid
        if not valid:
            print(f"    vt: {row['company_id']} failed thresholds")
    failures.append(not check("Top 5 verify ROE/FCF/RevCAGR constraints", ok))
    print("    top 5:", ", ".join(qc["company_id"].tolist()))

    composite_sorted = qc["composite_quality_score"].tolist() == sorted(
        qc["composite_quality_score"].tolist(), reverse=True
    )
    failures.append(not check("Top 5 sorted by composite score desc", composite_sorted))

    # ------------------------------------------------------------
    # 3. Screener workbook
    # ------------------------------------------------------------
    print("\n[3] output/screener_output.xlsx")
    xlsx_path = ROOT / "output" / "screener_output.xlsx"
    wb = load_workbook(xlsx_path)
    expected_sheets = {"All Companies", "Summary"} | {
        config["presets"][k]["name"] for k in config["presets"]
    }
    has_preset_sheets = all(s in wb.sheetnames for s in expected_sheets)
    failures.append(
        not check(
            "8 sheets (All + 6 presets + Summary) present",
            has_preset_sheets,
            str(wb.sheetnames),
        )
    )

    summary = wb["Summary"]
    summary_ok = True
    for row in summary.iter_rows(min_row=2, values_only=True):
        if row[2] and "OUT OF RANGE" in str(row[2]):
            summary_ok = False
    failures.append(not check("Summary sheet: all presets 'OK'", summary_ok))

    ws = wb[config["presets"]["quality_compounder"]["name"]]
    had_green = had_red = False
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=23):
        for cell in r:
            fill = cell.fill
            if fill.fill_type == "solid":
                if fill.start_color.rgb == "00C6EFCE":
                    had_green = True
                elif fill.start_color.rgb == "00FFC7CE":
                    had_red = True
    failures.append(
        not check(
            "Threshold cells colour-coded (green + red seen)",
            had_green and had_red,
            f"green={had_green} red={had_red}",
        )
    )

    # ------------------------------------------------------------
    # 4. IT Services peer ROE ranking
    # ------------------------------------------------------------
    print("\n[4] IT Services peer ROE ranking")
    import sqlite3

    conn = sqlite3.connect(ROOT / "db" / "nifty100.db")
    it_roe = pd.read_sql(
        "SELECT company_id, percentile_rank FROM peer_percentiles "
        "WHERE peer_group_name = 'IT Services' AND metric = 'roe' "
        "ORDER BY percentile_rank DESC",
        conn,
    )
    conn.close()
    # TCS has the highest ROE in IT Services -> must be percentile 1.00.
    tcs_top = (
        "TCS" in it_roe["company_id"].tolist()
        and it_roe.iloc[0]["company_id"] == "TCS"
        and abs(it_roe.iloc[0]["percentile_rank"] - 1.0) < 1e-6
    )
    failures.append(not check("TCS tops IT Services ROE percentile", tcs_top))
    print("    ranking:", it_roe.to_string(index=False).replace("\n", " ; "))

    # ------------------------------------------------------------
    # 5. Peer percentiles table + messages
    # ------------------------------------------------------------
    print("\n[5] peer_percentiles table")
    conn = sqlite3.connect(ROOT / "db" / "nifty100.db")
    n_rows = conn.execute("SELECT COUNT(*) FROM peer_percentiles").fetchone()[0]
    n_cos = conn.execute(
        "SELECT COUNT(DISTINCT company_id) FROM peer_percentiles"
    ).fetchone()[0]
    conn.close()
    failures.append(
        not check("peer_percentiles populated", n_rows > 0, f"{n_rows} rows")
    )
    failures.append(
        not check("56 peer-group companies covered", n_cos == 56, f"{n_cos} companies")
    )

    from src.analytics.peer import peer_percentile_for

    msg = peer_percentile_for("ABB")
    failures.append(
        not check(
            "No-peer-group message",
            "No peer group assigned" in msg,
            str(msg),
        )
    )

    # ------------------------------------------------------------
    # 6. Peer comparison workbook
    # ------------------------------------------------------------
    print("\n[6] output/peer_comparison.xlsx")
    peer_xlsx = ROOT / "output" / "peer_comparison.xlsx"
    wb = load_workbook(peer_xlsx)
    failures.append(
        not check(
            "11 peer-group sheets",
            len(wb.sheetnames) == 11,
            str(wb.sheetnames),
        )
    )

    it_sheet = wb["IT Services"]
    has_benchmark = False
    for r in range(2, it_sheet.max_row + 1):
        fill = it_sheet.cell(row=r, column=1).fill
        if fill.fill_type == "solid" and fill.start_color.rgb == "00FFD700":
            if it_sheet.cell(row=r, column=1).value == "TCS":
                has_benchmark = True
    failures.append(not check("Benchmark row (TCS) highlighted gold", has_benchmark))

    # ------------------------------------------------------------
    # 7. Radar charts
    # ------------------------------------------------------------
    print("\n[7] reports/radar_charts/")
    radar_dir = ROOT / "reports" / "radar_charts"
    radars = list(radar_dir.glob("*_radar.png"))
    bars = list(radar_dir.glob("*_bar.png"))
    failures.append(
        not check(
            "56 radar + 36 bar charts",
            len(radars) == 56 and len(bars) == 36,
            f"{len(radars)} radar, {len(bars)} bar",
        )
    )

    # ------------------------------------------------------------
    # Summary
    # ------------------------------------------------------------
    print("\n" + "=" * 70)
    total_fails = sum(1 for f in failures if f)
    if total_fails == 0:
        print("ALL SPRINT 3 CHECKS PASSED")
        print("Preset counts:", counts)
    else:
        print(f"{total_fails} CHECK(S) FAILED")
    print("=" * 70)
    return total_fails


if __name__ == "__main__":
    raise SystemExit(main())
