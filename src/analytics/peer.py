"""Peer Analysis — Sprint 3, Days 18–20.

Percentile ranks within peer groups, radar charts, and comparison workbook.
"""

import sqlite3
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

DB_PATH = Path("db/nifty100.db")

PEER_METRICS = {
    "roe": "return_on_equity_pct",
    "roce": "return_on_capital_employed_pct",
    "npm": "net_profit_margin_pct",
    "de": "debt_to_equity",
    "fcf": "free_cash_flow_cr",
    "pat_cagr_5yr": "pat_cagr_5yr",
    "rev_cagr_5yr": "revenue_cagr_5yr",
    "icr": "interest_coverage",
    "asset_turnover": "asset_turnover",
    "composite": "composite_quality_score",
}

RADAR_KEYS = [
    "return_on_equity_pct",
    "return_on_capital_employed_pct",
    "net_profit_margin_pct",
    "debt_to_equity",
    "free_cash_flow_cr",
    "pat_cagr_5yr",
    "revenue_cagr_5yr",
    "composite_quality_score",
]

RADAR_LABELS = [
    "ROE",
    "ROCE",
    "NPM",
    "D/E\n(inv)",
    "FCF",
    "PAT\nCAGR 5y",
    "Rev\nCAGR 5y",
    "Composite",
]

KPI_COLUMNS = [
    "return_on_equity_pct",
    "return_on_capital_employed_pct",
    "net_profit_margin_pct",
    "operating_profit_margin_pct",
    "return_on_assets_pct",
    "debt_to_equity",
    "interest_coverage",
    "asset_turnover",
    "free_cash_flow_cr",
    "cash_from_operations_cr",
    "total_debt_cr",
    "earnings_per_share",
    "book_value_per_share",
    "dividend_payout_ratio_pct",
    "revenue_cagr_5yr",
    "pat_cagr_5yr",
    "eps_cagr_5yr",
    "pe_ratio",
    "pb_ratio",
    "dividend_yield_pct",
]


# ------------------------------------------------------------------
# Data loaders
# ------------------------------------------------------------------


def _load_peer_groups(db_path: Path = DB_PATH) -> pd.DataFrame:
    """Load peer groups from the raw xlsx (Sprint 3 spec)."""
    try:
        xlsx = next(p for p in Path("data/supporting").glob("*peer_groups.xlsx"))
        df = pd.read_excel(xlsx, dtype={"company_id": str, "is_benchmark": bool})
    except (StopIteration, FileNotFoundError):
        conn = sqlite3.connect(db_path)
        df = pd.read_sql(
            "SELECT peer_group_name, company_id, is_benchmark FROM peer_groups",
            conn,
        )
        conn.close()
        df["is_benchmark"] = df["is_benchmark"].map(
            {"1": True, "0": False, True: True, False: False}
        )
    return df[["peer_group_name", "company_id", "is_benchmark"]]


def _load_baseline_frame(db_path: Path = DB_PATH) -> pd.DataFrame:
    """Load the full 92-company baseline feature frame."""
    from src.screener.engine import compute_composite_score, load_feature_frame

    return compute_composite_score(load_feature_frame(db_path=db_path))


# ------------------------------------------------------------------
# Day 18 — Percentile ranks
# ------------------------------------------------------------------


def _percent_rank(values: pd.Series, ascending: bool = False) -> pd.Series:
    """SQL-style PERCENT_RANK: (rank − 1) / (n − 1), ties = min rank."""

    valid = values.dropna()
    if len(valid) < 2:
        return pd.Series(np.nan, index=values.index)

    n = len(valid)
    ranks = valid.rank(method="min", ascending=ascending)
    pct = (ranks - 1) / (n - 1)

    result = pd.Series(np.nan, index=values.index)
    result.loc[pct.index] = pct
    return result


def compute_peer_percentiles(
    frame: pd.DataFrame, peer_groups: pd.DataFrame
) -> pd.DataFrame:
    """Compute PERCENT_RANK for every metric within each peer group.

    D/E is inverted so that lower D/E = higher (better) percentile.
    Returns a DataFrame with columns:
        company_id, peer_group_name, metric, value, percentile_rank, year
    """

    rows = []

    for group_name, group_df in peer_groups.groupby("peer_group_name"):
        ids = set(group_df["company_id"])
        group_frame = frame[frame["company_id"].isin(ids)].copy()

        if len(group_frame) < 2:
            continue

        year_val = (
            str(group_frame["year"].mode().iloc[0])
            if "year" in group_frame.columns
            else None
        )

        for metric_key, col in PEER_METRICS.items():
            if col not in group_frame.columns:
                continue

            valid = group_frame[["company_id", col]].dropna(subset=[col])
            if len(valid) < 2:
                continue

            pct = _percent_rank(valid[col], ascending=True)

            for i, p in pct.items():
                if pd.isna(p):
                    continue
                cid = valid.loc[i, "company_id"]
                value = valid.loc[i, col]
                if metric_key == "de":
                    p = 1 - p
                rows.append(
                    {
                        "company_id": cid,
                        "peer_group_name": group_name,
                        "metric": metric_key,
                        "value": value,
                        "percentile_rank": round(float(p), 4),
                        "year": year_val,
                    }
                )

    return pd.DataFrame(rows)


def persist_peer_percentiles(
    percentiles: pd.DataFrame, db_path: Path = DB_PATH
) -> None:
    """Write peer_percentiles to SQLite (idempotent replace)."""

    conn = sqlite3.connect(db_path)
    conn.execute("DROP TABLE IF EXISTS peer_percentiles")
    conn.execute("""
        CREATE TABLE IF NOT EXISTS peer_percentiles (
            id            INTEGER PRIMARY KEY,
            company_id    TEXT NOT NULL,
            peer_group_name TEXT NOT NULL,
            metric        TEXT NOT NULL,
            value         NUMERIC,
            percentile_rank NUMERIC,
            year          TEXT,
            FOREIGN KEY (company_id) REFERENCES companies(id)
                ON UPDATE CASCADE ON DELETE RESTRICT,
            UNIQUE (company_id, peer_group_name, metric)
        )
        """)
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_pp_group_metric "
        "ON peer_percentiles(peer_group_name, metric)"
    )

    for _, row in percentiles.iterrows():
        conn.execute(
            "INSERT OR REPLACE INTO peer_percentiles "
            "(company_id, peer_group_name, metric, value, percentile_rank, year) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (
                row["company_id"],
                row["peer_group_name"],
                row["metric"],
                row["value"],
                row["percentile_rank"],
                row["year"],
            ),
        )

    conn.commit()
    conn.close()


def peer_percentile_for(company_id: str, db_path: Path = DB_PATH) -> str:
    """Return percentile ranks for a company, or a 'no peer group' message."""

    conn = sqlite3.connect(db_path)
    try:
        table_exists = conn.execute(
            "SELECT name FROM sqlite_master "
            "WHERE type = 'table' AND name = 'peer_percentiles'"
        ).fetchone()
        if not table_exists:
            return f"No peer group assigned for {company_id}"

        df = pd.read_sql(
            "SELECT metric, peer_group_name, percentile_rank "
            "FROM peer_percentiles WHERE company_id = ?",
            conn,
            params=[company_id],
        )
    finally:
        conn.close()

    if df.empty:
        return f"No peer group assigned for {company_id}"

    return df.to_string(index=False)


# ------------------------------------------------------------------
# Day 19 — Radar / bar charts
# ------------------------------------------------------------------


def _normalise_peer_scores(
    group_frame: pd.DataFrame,
) -> pd.DataFrame:
    """Winsorised 0–100 scores per axis for a peer group.

    Returns a DataFrame indexed by company_id with columns aligned
    to RADAR_KEYS. D/E is inverted (lower = better).
    """

    scores = {}

    for col in RADAR_KEYS:
        s = group_frame[col].dropna()
        if len(s) < 2:
            scores[col] = pd.Series(50.0, index=group_frame.index)
            continue

        p10, p90 = s.quantile(0.10), s.quantile(0.90)
        if p10 == p90:
            scores[col] = pd.Series(50.0, index=group_frame.index)
            continue

        clipped = s.clip(p10, p90)
        normed = ((clipped - p10) / (p90 - p10)) * 100

        if col == "debt_to_equity":
            normed = 100 - normed

        full = pd.Series(np.nan, index=group_frame.index)
        full.loc[normed.index] = normed
        scores[col] = full

    result = pd.DataFrame(scores, index=group_frame.index)
    result.index = group_frame["company_id"].values
    result.index.name = "company_id"
    return result


def _plot_radar(
    company_id: str,
    company_scores: np.ndarray,
    avg_scores: np.ndarray,
    group_name: str,
    output_dir: Path,
) -> None:

    n = len(RADAR_LABELS)
    angles = np.linspace(0, 2 * np.pi, n, endpoint=False).tolist()
    company_plot = np.append(company_scores, company_scores[0])
    avg_plot = np.append(avg_scores, avg_scores[0])
    angles_plot = angles + [angles[0]]

    fig, ax = plt.subplots(figsize=(8, 8), subplot_kw=dict(polar=True))
    ax.fill(angles_plot, company_plot, alpha=0.25, color="#1f77b4")
    ax.plot(angles_plot, company_plot, color="#1f77b4", linewidth=2, label=company_id)
    ax.plot(
        angles_plot,
        avg_plot,
        color="#ff7f0e",
        linewidth=1.5,
        linestyle="--",
        label=f"{group_name} avg",
    )
    ax.set_thetagrids(np.degrees(angles), RADAR_LABELS)
    ax.set_ylim(0, 100)
    ax.set_title(f"{company_id}  vs  {group_name}", y=1.08, size=14)
    ax.legend(loc="upper right", bbox_to_anchor=(1.35, 1.12), fontsize=9)
    fig.tight_layout()
    fig.savefig(output_dir / f"{company_id}_radar.png", dpi=150, bbox_inches="tight")
    plt.close(fig)


def _plot_bar(
    company_id: str,
    company_row: pd.Series,
    nifty_avg: pd.Series,
    output_dir: Path,
) -> None:

    x = np.arange(len(RADAR_LABELS))
    company_vals = [
        float(company_row.get(c, 0)) if pd.notna(company_row.get(c)) else 0
        for c in RADAR_KEYS
    ]
    avg_vals = [
        float(nifty_avg.get(c, 0)) if pd.notna(nifty_avg.get(c)) else 0
        for c in RADAR_KEYS
    ]

    fig, ax = plt.subplots(figsize=(10, 6))
    ax.bar(x - 0.15, company_vals, 0.3, label=company_id, color="#1f77b4")
    ax.bar(x + 0.15, avg_vals, 0.3, label="Nifty 100 avg", color="#ff7f0e")
    ax.set_xticks(x)
    ax.set_xticklabels(RADAR_LABELS, rotation=45, ha="right")
    ax.set_ylabel("Raw value")
    ax.set_title(f"{company_id} — No peer group (vs Nifty 100 average)")
    ax.legend()
    fig.tight_layout()
    fig.savefig(output_dir / f"{company_id}_bar.png", dpi=150, bbox_inches="tight")
    plt.close(fig)


def generate_radar_charts(
    frame: pd.DataFrame,
    peer_groups: pd.DataFrame,
    output_dir: Path = Path("reports/radar_charts"),
) -> list[Path]:
    """Generate one chart per company.

    - Companies in a peer group: polar radar with peer-average overlay.
    - Companies with no peer group: bar chart vs Nifty 100 average.

    Returns list of saved file paths.
    """

    output_dir.mkdir(parents=True, exist_ok=True)
    nifty_avg = frame[RADAR_KEYS].mean()
    group_lookup = peer_groups.set_index("company_id")["peer_group_name"].to_dict()
    saved: list[Path] = []

    for group_name, group_df in peer_groups.groupby("peer_group_name"):
        ids = group_df["company_id"].tolist()
        group_frame = frame[frame["company_id"].isin(ids)].copy()

        if len(group_frame) < 2:
            continue

        normed = _normalise_peer_scores(group_frame)
        avg_scores = normed.mean().values

        for cid in ids:
            if cid not in normed.index:
                continue
            company_scores = normed.loc[cid].values
            _plot_radar(cid, company_scores, avg_scores, group_name, output_dir)
            saved.append(output_dir / f"{cid}_radar.png")

    no_group = set(frame["company_id"]) - set(group_lookup)
    for cid in sorted(no_group):
        company_row = frame[frame["company_id"] == cid]
        if company_row.empty:
            continue
        _plot_bar(cid, company_row.iloc[0], nifty_avg, output_dir)
        saved.append(output_dir / f"{cid}_bar.png")

    return saved


# ------------------------------------------------------------------
# Day 20 — Peer comparison workbook
# ------------------------------------------------------------------


def _pct_fill(rank: float):
    """Return openpyxl PatternFill colour for a percentile rank."""

    from openpyxl.styles import PatternFill

    if rank >= 0.75:
        return PatternFill(
            start_color="00C6EFCE", end_color="00C6EFCE", fill_type="solid"
        )
    if rank <= 0.25:
        return PatternFill(
            start_color="00FFC7CE", end_color="00FFC7CE", fill_type="solid"
        )
    return PatternFill(start_color="00FFEB9C", end_color="00FFEB9C", fill_type="solid")


def write_peer_comparison_xlsx(
    frame: pd.DataFrame,
    peer_groups: pd.DataFrame,
    percentiles: pd.DataFrame,
    output_path: Path = Path("output/peer_comparison.xlsx"),
) -> Path:
    """Generate the 11-sheet peer comparison workbook.

    Each sheet: company_id, company_name, 20 KPI columns,
    10 percentile columns (green / yellow / red), benchmark
    rows with gold background, and a median summary row.
    """

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)

    benchmark_fill = PatternFill(
        start_color="00FFD700", end_color="00FFD700", fill_type="solid"
    )
    bold = Font(bold=True)

    pct_metric_keys = list(PEER_METRICS.keys())
    pct_col_headers = [f"pct_{k}" for k in pct_metric_keys]

    for group_name, group_df in peer_groups.groupby("peer_group_name"):
        sheet_title = group_name[:31]
        ws = wb.create_sheet(sheet_title)

        ids = group_df["company_id"].tolist()
        benchmarks = set(group_df.loc[group_df["is_benchmark"], "company_id"])
        group_frame = frame[frame["company_id"].isin(ids)].copy()
        group_frame = group_frame.sort_values(
            "composite_quality_score", ascending=False
        )

        headers = ["company_id", "company_name"] + KPI_COLUMNS + pct_col_headers
        for ci, h in enumerate(headers, 1):
            ws.cell(row=1, column=ci, value=h).font = bold

        for ri, (_, row) in enumerate(group_frame.iterrows(), 2):
            ws.cell(row=ri, column=1, value=row.get("company_id"))
            ws.cell(row=ri, column=2, value=row.get("company_name"))

            for ci, col in enumerate(KPI_COLUMNS, 3):
                val = row.get(col)
                ws.cell(row=ri, column=ci, value=val)

            for pi, mk in enumerate(pct_metric_keys, 3 + len(KPI_COLUMNS)):
                match = percentiles[
                    (percentiles["company_id"] == row.get("company_id"))
                    & (percentiles["peer_group_name"] == group_name)
                    & (percentiles["metric"] == mk)
                ]
                if not match.empty:
                    pr = float(match.iloc[0]["percentile_rank"])
                    cell = ws.cell(row=ri, column=pi, value=round(pr, 4))
                    cell.fill = _pct_fill(pr)

            if row.get("company_id") in benchmarks:
                for ci in range(1, ws.max_column + 1):
                    ws.cell(row=ri, column=ci).fill = benchmark_fill

        summary_row = ws.max_row + 2
        ws.cell(row=summary_row, column=1, value="Median").font = bold

        for ci, col in enumerate(KPI_COLUMNS, 3):
            med = group_frame[col].median()
            ws.cell(
                row=summary_row,
                column=ci,
                value=round(float(med), 2) if pd.notna(med) else None,
            )

        for pi, mk in enumerate(pct_metric_keys, 3 + len(KPI_COLUMNS)):
            med = percentiles[
                (percentiles["peer_group_name"] == group_name)
                & (percentiles["metric"] == mk)
            ]["percentile_rank"].median()
            ws.cell(
                row=summary_row,
                column=pi,
                value=round(float(med), 4) if pd.notna(med) else None,
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


# ------------------------------------------------------------------
# Main (Day 18–20 pipeline)
# ------------------------------------------------------------------


def main() -> None:
    """Full peer-analysis pipeline: percentiles → radar → xlsx."""

    print("=" * 60)
    print("PEER ANALYSIS  (Sprint 3 — Days 18–20)")
    print("=" * 60)

    frame = _load_baseline_frame()
    peer_groups = _load_peer_groups()

    print(f"\n  Companies in feature frame : {len(frame)}")
    print(f"  Peer groups                : {peer_groups['peer_group_name'].nunique()}")
    print(f"  Companies in peer groups   : {peer_groups['company_id'].nunique()}")

    print("\n[1/3] Computing percentile ranks …")
    percentiles = compute_peer_percentiles(frame, peer_groups)
    persist_peer_percentiles(percentiles)
    print(f"  {len(percentiles)} rows written to peer_percentiles")

    print("\n[2/3] Generating radar / bar charts …")
    charts = generate_radar_charts(frame, peer_groups)
    print(f"  {len(charts)} charts saved to reports/radar_charts/")

    print("\n[3/3] Writing peer comparison workbook …")
    xlsx_path = write_peer_comparison_xlsx(frame, peer_groups, percentiles)
    print(f"  {xlsx_path}")

    print("\n" + "=" * 60)
    print("PEER ANALYSIS COMPLETE")
    print("=" * 60)


if __name__ == "__main__":
    main()
